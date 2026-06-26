import logging
import openpyxl
import csv
from django.http import HttpResponse
from .models import Product, Region, Lead, ProductCategory,LeadStatus, LeadSource, Territory
from .forms import ProductForm, RegionForm, LeadForm, ProductBulkUploadForm, LeadBulkUploadForm
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.db.models import Max
from django.db import transaction
from django.db.models import Q
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.decorators import api_view, permission_classes
from django.contrib import messages
from .serializers import (ProductSerializer,RegionSerializer,LeadSerializer)
from django.contrib.auth.decorators import login_required
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

logger = logging.getLogger(__name__)
@login_required
def home(request):
    return render(request, 'home.html')

@login_required
def product_list(request):
    query = request.GET.get('q','')
    products = Product.objects.all()
    if query:
        products = products.filter(
            Q(productid__icontains=query) |
            Q(productname__icontains=query) |
            Q(categoryid__categoryname__icontains=query) |
            Q(added_by__icontains=query) |
            Q(added_dts__icontains=query)
        )

    return render(request, 'product/product_list.html', {
        'products': products,
        'query': query
    })

@login_required
def lead_export_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="leads.csv"'

    writer = csv.writer(response)

    writer.writerow([
        'Lead ID', 'Person Name', 'Gender', 'Contact No', 'Email',
        'Company', 'City', 'State', 'Territory', 'Region', 'Product',
        'Status', 'Lead Source', 'Business Need',
        'Lead Generated Date', 'Added By', 'Added DTS'
    ])

    leads = Lead.objects.select_related(
        'territoryid', 'regionid', 'productid', 'statusid', 'leadsourceid'
    )

    for lead in leads:
        writer.writerow([
            lead.leadid,
            lead.personname,
            lead.gender,
            lead.contactno,
            lead.email,
            lead.companyname,
            lead.city,
            lead.state,
            lead.territoryid.territoryname if lead.territoryid else '',
            lead.regionid.regionname if lead.regionid else '',
            lead.productid.productname if lead.productid else '',
            lead.statusid.statusname if lead.statusid else '',
            lead.leadsourceid.leadsourcename if lead.leadsourceid else '',
            lead.businessneed,
            lead.lead_gen_date.strftime('%d-%m-%Y') if lead.lead_gen_date else '',
            lead.added_by,
            lead.added_dts.strftime('%d-%m-%Y %H:%M') if lead.added_dts else ''
        ])

    return response

@login_required
def product_export_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="products.csv"'

    writer = csv.writer(response)

    # Header row
    writer.writerow([
        'Product ID',
        'Product Name',
        'Category ID',
        'Is Active',
        'Added By',
        'Added Date'
    ])

    # Data rows
    products = Product.objects.all()

    for p in products:
        writer.writerow([
            p.productid,
            p.productname,
            p.categoryid.categoryid if p.categoryid else '',
            p.is_active,
            p.added_by,
            p.added_dts.strftime("%Y-%m-%d %H:%M:%S") if p.added_dts else ''
        ])

    return response

def region_export_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="regions.csv"'

    writer = csv.writer(response)

    writer.writerow([
        'Region ID',
        'Region Name',
        'Added By',
        'Added Date'
    ])

    regions = Region.objects.all()

    for r in regions:
        writer.writerow([
            r.regionid,
            r.regionname,
            r.added_by,
            r.added_dts.strftime("%Y-%m-%d %H:%M:%S") if r.added_dts else ''
        ])

    return response


@login_required
def lead_export_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    # ---------- HEADER ROW ----------
    headers = [
        'Lead ID', 'Person Name', 'Gender', 'Contact No', 'Email',
        'Company', 'City', 'State', 'Territory', 'Region', 'Product',
        'Status', 'Lead Source', 'Business Need',
        'Lead Generated Date', 'Added By', 'Added DTS'
    ]

    ws.append(headers)

    # Header styling
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')

    # ---------- DATA ----------
    leads = Lead.objects.select_related(
        'territoryid', 'regionid', 'productid', 'statusid', 'leadsourceid'
    )

    for lead in leads:
        ws.append([
            lead.leadid,
            lead.personname,
            lead.gender,
            lead.contactno,
            lead.email,
            lead.companyname,
            lead.city,
            lead.state,
            lead.territoryid.territoryname if lead.territoryid else '',
            lead.regionid.regionname if lead.regionid else '',
            lead.productid.productname if lead.productid else '',
            lead.statusid.statusname if lead.statusid else '',
            lead.leadsourceid.leadsourcename if lead.leadsourceid else '',
            lead.businessneed,
            lead.lead_gen_date.strftime('%d-%m-%Y') if lead.lead_gen_date else '',
            lead.added_by,
            lead.added_dts.strftime('%d-%m-%Y %H:%M') if lead.added_dts else ''
        ])

    # ---------- AUTO COLUMN WIDTH ----------
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 2

    # ---------- RESPONSE ----------
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=leads.xlsx'

    wb.save(response)
    return response

@login_required
def add_product(request):
    try:
        if request.method == 'POST':
            form = ProductForm(request.POST)
            if form.is_valid():
                product = form.save(commit=False)
                max_id = Product.objects.aggregate(Max('productid'))['productid__max']
                product.productid = (max_id or 0) + 1
                product.added_by = request.user.username
                product.added_dts = timezone.now()
                product.save()
                return redirect('product_list')
        else:
            form = ProductForm()

        return render(request, 'product/product_form.html', {'form': form, 'title': 'Add Product'})

    except Exception:
        logger.exception("Unexpected error in add_product")
        messages.error(request, "Something went wrong while adding the product. Please try again.")
        return redirect('product_list')

@login_required
def edit_product(request, id):
    try:
        product = get_object_or_404(Product, pk=id)

        if request.method == 'POST':
            form = ProductForm(request.POST, instance=product)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.added_by = request.user.username
                obj.added_dts = timezone.now()
                obj.save()
                return redirect('product_list')
        else:
            form = ProductForm(instance=product)

        return render(request, 'product/product_form.html', {'form': form, 'title': 'Edit Product'})

    except Exception:
        logger.exception("Unexpected error in edit_product")
        return redirect('product_list')

@login_required
def delete_product(request, id):
    try:
        product = get_object_or_404(Product, pk=id)
        product.delete()
        return redirect('product_list')
    except Exception:
        logger.exception("Unexpected error in delete_product")
        return redirect('product_list')
@login_required
def product_bulk_upload(request):
    if request.method == 'POST':
        form = ProductBulkUploadForm(request.POST, request.FILES)

        if form.is_valid():
            try:
                wb = openpyxl.load_workbook(request.FILES['file'])
                sheet = wb.active

                rows = list(sheet.iter_rows(min_row=2, values_only=True))

                if len(rows) == 0:
                    messages.error(request, "Excel file does not contain any products.")
                    return redirect('product_bulk_upload')

                # optional but VERY GOOD practice
                with transaction.atomic():

                    for index, row in enumerate(rows, start=2):
                        product_name, category_name, isactive = row

                        # Validation 1: product name
                        if not product_name:
                            raise ValueError(f"Row {index}: Product name cannot be empty")

                        # Validation 2: category
                        if not category_name:
                            raise ValueError(f"Row {index}: Category cannot be empty")

                        # Validation 3: duplicate product
                        if Product.objects.filter(productname=product_name).exists():
                            raise ValueError(f"Row {index}: Duplicate product '{product_name}'")

                        # Convert category string → ProductCategory object
                        try:
                            category_obj = ProductCategory.objects.get(
                                categoryname__iexact=str(category_name).strip()
                            )
                        except ProductCategory.DoesNotExist:
                            raise ValueError(
                                f"Row {index}: Category '{category_name}' does not exist"
                            )

                        # Default is_active
                        if isactive is None:
                            isactive = 1

                        max_id = Product.objects.aggregate(
                            Max('productid')
                        )['productid__max']

                        Product.objects.create(
                            productid=(max_id or 0) + 1,
                            productname=product_name,
                            categoryid=category_obj,   
                            is_active=int(isactive),
                            added_by=request.user.username,
                            added_dts=timezone.now()
                        )

                messages.success(request, "Products uploaded successfully")
                return redirect('product_list')

            except Exception as e:
                logger.exception("Bulk upload failed")
                messages.error(request, str(e))
                return redirect('product_bulk_upload')

    else:
        form = ProductBulkUploadForm()

    return render(request, 'product/product_bulk_upload.html', {'form': form})

@login_required
def region_list(request):
    query = request.GET.get('q','')
    regions = Region.objects.all()

    if query:
        regions = regions.filter(
            Q(regionid__icontains=query) |
            Q(regionname__icontains=query) |
            Q(added_by__icontains=query) |
            Q(added_dts__icontains=query)
        )

    return render(request, 'region/region_list.html', {
        'regions': regions,
        'query': query
    })

@login_required
def add_region(request):
    try:
        if request.method == 'POST':
            form = RegionForm(request.POST)
            if form.is_valid():
                region = form.save(commit=False)
                max_id = Region.objects.aggregate(Max('regionid'))['regionid__max']
                region.regionid = (max_id or 0) + 1
                region.added_by = request.user.username
                region.added_dts = timezone.now()
                region.save()
                return redirect('region_list')
        else:
            form = RegionForm()

        return render(request, 'region/region_form.html', {'form': form, 'title': 'Add Region'})

    except Exception:
        logger.exception("Unexpected error in add_region")
        return redirect('region_list')

@login_required
def edit_region(request, id):
    try:
        region = get_object_or_404(Region, pk=id)

        if request.method == 'POST':
            form = RegionForm(request.POST, instance=region)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.added_by = request.user.username
                obj.added_dts = timezone.now()
                obj.save()
                return redirect('region_list')
        else:
            form = RegionForm(instance=region)

        return render(request, 'region/region_form.html', {'form': form, 'title': 'Edit Region'})

    except Exception:
        logger.exception("Unexpected error in edit_region")
        return redirect('region_list')

@login_required
def delete_region(request, id):
    try:
        region = get_object_or_404(Region, pk=id)
        region.delete()
        return redirect('region_list')
    except Exception:
        logger.exception("Unexpected error in delete_region")
        return redirect('region_list')

@login_required
def lead_list(request):
    query = request.GET.get('q','')
    leads = Lead.objects.all()

    if query:
        leads = leads.filter(
            Q(leadid__icontains=query) |
            Q(personname__icontains=query) |
            Q(companyname__icontains=query) |
            Q(email__icontains=query) |
            Q(contactno__icontains=query) |
            Q(city__icontains=query) |
            Q(state__icontains=query) |
            Q(regionid__regionname__icontains=query) |
            Q(productid__productname__icontains=query) |
            Q(statusid__statusname__icontains=query) |
            Q(leadsourceid__leadsourcename__icontains=query) |
            Q(added_by__icontains=query) |
            Q(lead_gen_date__icontains=query)
        )

    return render(request, 'lead/lead_list.html', {
        'leads': leads,
        'query': query
    })

@login_required
def add_lead(request):
    try:
        if request.method == 'POST':
            form = LeadForm(request.POST)
            if form.is_valid():
                lead = form.save(commit=False)
                max_id = Lead.objects.aggregate(Max('leadid'))['leadid__max']
                lead.leadid = (max_id or 0) + 1
                lead.added_by = request.user.username
                lead.added_dts = timezone.now()
                lead.save()
                return redirect('lead_list')
        else:
            form = LeadForm()

        return render(request, 'lead/lead_form.html', {'form': form, 'title': 'Add Lead'})

    except Exception:
        logger.exception("Unexpected error in add_lead")
        return redirect('lead_list')
    
@login_required
def lead_bulk_upload(request):
    if request.method == 'POST':
        form = LeadBulkUploadForm(request.POST, request.FILES)

        if form.is_valid():
            try:
                wb = openpyxl.load_workbook(request.FILES['file'])
                sheet = wb.active
                rows = list(sheet.iter_rows(min_row=2, values_only=True))

                if not rows:
                    messages.error(request, "Excel file is empty.")
                    return redirect('lead_bulk_upload')

                for row in rows:
                    (
                        personname,gender,companyname,contactno,email,
                        city,state,territory_name,region_name,product_name,status_name,
                        leadsource_name,businessneed,lead_gen_date,executiveid
                    ) = row

                    # -------- BASIC VALIDATIONS --------
                    if not personname:
                        raise ValueError("Person name is required")

                    if not email:
                        raise ValueError(f"Email missing for {personname}")

                    if Lead.objects.filter(email=email).exists():
                        raise ValueError(f"Duplicate email found: {email}")

                    # -------- REGION --------
                    try:
                        region_obj = Region.objects.get(
                            regionname__iexact=region_name
                        )
                    except Region.DoesNotExist:
                        raise ValueError(f"Region '{region_name}' does not exist")
                    
                    try:
                       territory_obj = Territory.objects.get(
                       territoryname__iexact=territory_name
                       )
                    except Territory.DoesNotExist:
                       raise ValueError(f"Territory '{territory_name}' does not exist")

                    # -------- PRODUCT --------
                    try:
                        product_obj = Product.objects.get(
                            productname__iexact=product_name
                        )
                    except Product.DoesNotExist:
                        raise ValueError(f"Product '{product_name}' does not exist")

                    # -------- STATUS --------
                    try:
                        status_obj = LeadStatus.objects.get(
                            statusname__iexact=status_name
                        )
                    except LeadStatus.DoesNotExist:
                        raise ValueError(f"Status '{status_name}' does not exist")

                    # -------- LEAD SOURCE --------
                    try:
                        source_obj = LeadSource.objects.get(
                            leadsourcename__iexact=leadsource_name
                        )
                    except LeadSource.DoesNotExist:
                        raise ValueError(f"Lead Source '{leadsource_name}' does not exist")

                    max_id = Lead.objects.aggregate(
                        Max('leadid')
                    )['leadid__max']
                    new_lead_id = (max_id or 0)+1

                    Lead.objects.create(
                        leadid = new_lead_id,
                        personname=personname,
                        gender=gender,
                        companyname=companyname,
                        contactno=contactno,
                        email=email,
                        city=city,
                        state=state,
                        territoryid=territory_obj,
                        regionid=region_obj,
                        productid=product_obj,
                        statusid=status_obj,
                        leadsourceid=source_obj,
                        businessneed=businessneed,
                        lead_gen_date=lead_gen_date,
                        added_by = request.user.username,
                        added_dts = timezone.now(),
                        executiveid=executiveid
                    )

                messages.success(request, "Leads uploaded successfully")
                return redirect('lead_list')

            except Exception as e:
                logger.exception("Lead bulk upload failed")
                messages.error(request, str(e))
                return redirect('lead_bulk_upload')

    else:
        form = LeadBulkUploadForm()

    return render(request, 'lead/lead_bulk_upload.html', {'form': form})

@login_required
def edit_lead(request, id):
    try:
        lead = get_object_or_404(Lead, pk=id)

        if request.method == 'POST':
            form = LeadForm(request.POST, instance=lead)
            if form.is_valid():
                obj = form.save(commit = False)
                obj.added_by = request.user.username
                obj.added_dts = timezone.now()
                obj.save()
                return redirect('lead_list')
        else:
            form = LeadForm(instance=lead)

        return render(request, 'lead/lead_form.html', {'form': form, 'title': 'Edit Lead'})

    except Exception:
        logger.exception("Unexpected error in edit_lead")
        return redirect('lead_list')

@login_required
def delete_lead(request, id):
    try:
        lead = get_object_or_404(Lead, pk=id)
        lead.delete()
        return redirect('lead_list')
    except Exception:
        logger.exception("Unexpected error in delete_lead")
        return redirect('lead_list')

@api_view(['GET'])
def product_api(request):

    products = Product.objects.all()

    serializer = ProductSerializer(
        products,
        many=True
    )

    return Response({
        "success": True,
        "count": products.count(),
        "data": serializer.data
    })

@api_view(['GET'])
def product_detail_api(request, productid):
    try:
        product = Product.objects.filter(pk=productid).first()

        if not product:
            return Response(
                {"success": False, "message": "Product Not Found"},
                status=404
            )

        serializer = ProductSerializer(product)
        return Response({"success": True, "data": serializer.data})

    except Exception:
        logger.exception("Unexpected error in product_detail_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['GET'])
def region_api(request):

    regions = Region.objects.all()

    serializer = RegionSerializer(
        regions,
        many=True
    )

    return Response({
        "success": True,
        "count": regions.count(),
        "data": serializer.data
    })

api_view(['GET'])
def region_detail_api(request, regionid):
    try:
        region = Region.objects.filter(pk=regionid).first()

        if not region:
            return Response(
                {"success": False, "message": "Region Not Found"},
                status=404
            )

        serializer = RegionSerializer(region)
        return Response({"success": True, "data": serializer.data})

    except Exception:
        logger.exception("Unexpected error in region_detail_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['GET'])
def lead_api(request):

    leads = Lead.objects.all()

    serializer = LeadSerializer(
        leads,
        many=True
    )

    return Response({
        "success": True,
        "count": leads.count(),
        "data": serializer.data
    })

@api_view(['GET'])
def lead_detail_api(request, leadid):
    try:
        lead = Lead.objects.filter(pk=leadid).first()

        if not lead:
            return Response(
                {"success": False, "message": "Lead Not Found"},
                status=404
            )

        serializer = LeadSerializer(lead)
        return Response({"success": True, "data": serializer.data})

    except Exception:
        logger.exception("Unexpected error in lead_detail_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def product_create_api(request):
    try:
        serializer = ProductSerializer(data=request.data)

        if not serializer.is_valid():
            return Response(
                {"success": False, "errors": serializer.errors},
                status=400
            )

        max_id = Product.objects.aggregate(
            Max('productid')
        )['productid__max']

        serializer.save(
            productid=(max_id or 0) + 1,
            added_by=request.user.username,
            added_dts=timezone.now()
        )

        return Response(
            {"success": True, "message": "Product added successfully"},
            status=201
        )

    except Exception as e:
        logger.exception(
            f"Unexpected error in product_create_api: {str(e)}"
        )
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def region_create_api(request):
    try:
        serializer = RegionSerializer(data=request.data)

        if not serializer.is_valid():
            return Response(
                {"success": False, "errors": serializer.errors},
                status=400
            )

        max_id = Region.objects.aggregate(
            Max('regionid')
        )['regionid__max']

        serializer.save(
            regionid=(max_id or 0) + 1,
            added_by=request.user.username,
            added_dts=timezone.now()
        )

        return Response(
            {"success": True, "message": "Region Added Successfully"},
            status=201
        )

    except Exception as e:
        logger.exception("Unexpected error in region_create_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def lead_create_api(request):
    try:
        serializer = LeadSerializer(data=request.data)

        if not serializer.is_valid():
            return Response(
                {"success": False, "errors": serializer.errors},
                status=400
            )

        max_id = Lead.objects.aggregate(
            Max('leadid')
        )['leadid__max']

        serializer.save(
            leadid=(max_id or 0) + 1,
            added_by=request.user.username,
            added_dts=timezone.now()
        )

        return Response(
            {"success": True, "message": "Lead Added Successfully"},
            status=201
        )

    except Exception:
        logger.exception("Unexpected error in lead_create_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def product_update_api(request, productid):
    try:
        product = Product.objects.filter(pk=productid).first()
        if not product:
            return Response({"success": False, "message": "Product Not Found"}, status=404)

        old_data = ProductSerializer(product).data
        serializer = ProductSerializer(product, data=request.data)

        if not serializer.is_valid():
            return Response({"success": False, "errors": serializer.errors}, status=400)

        serializer.save(added_by=request.user.username, added_dts=timezone.now())

        return Response({
            "success": True,
            "message": "Product Updated Successfully",
            "old_data": old_data,
            "new_data": serializer.data
        })

    except Exception:
        logger.exception("Unexpected error in product_update_api")
        return Response({"success": False, "message": "Internal server error"}, status=500)


@api_view(['DELETE'])
def product_delete_api(request, productid):
    try:
        product = Product.objects.filter(pk=productid).first()
        if not product:
            return Response({"success": False, "message": "Product Not Found"}, status=404)

        product.delete()
        return Response({"success": True, "message": "Product Deleted Successfully"})

    except Exception:
        logger.exception("Unexpected error in product_delete_api")
        return Response({"success": False, "message": "Internal server error"}, status=500)


@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def region_update_api(request, regionid):
    try:
        region = Region.objects.filter(pk=regionid).first()

        if not region:
            return Response(
                {"success": False, "message": "Region Not Found"},
                status=404
            )

        old_data = RegionSerializer(region).data
        serializer = RegionSerializer(region, data=request.data)

        if not serializer.is_valid():
            return Response(
                {"success": False, "errors": serializer.errors},
                status=400
            )

        serializer.save(
            added_by=request.user.username,
            added_dts=timezone.now()
        )

        return Response({
            "success": True,
            "message": "Region Updated Successfully",
            "old_data": old_data,
            "new_data": serializer.data
        })

    except Exception:
        logger.exception("Unexpected error in region_update_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )


@api_view(['DELETE'])
def region_delete_api(request, regionid):
    try:
        region = Region.objects.filter(pk=regionid).first()

        if not region:
            return Response(
                {"success": False, "message": "Region Not Found"},
                status=404
            )

        region.delete()
        return Response(
            {"success": True, "message": "Region Deleted Successfully"}
        )

    except Exception:
        logger.exception("Unexpected error in region_delete_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )

@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def lead_update_api(request, leadid):
    try:
        lead = Lead.objects.filter(pk=leadid).first()

        if not lead:
            return Response(
                {"success": False, "message": "Lead Not Found"},
                status=404
            )

        old_data = LeadSerializer(lead).data
        serializer = LeadSerializer(lead, data=request.data)

        if not serializer.is_valid():
            return Response(
                {"success": False, "errors": serializer.errors},
                status=400
            )

        serializer.save(
            added_by=request.user.username,
            added_dts=timezone.now()
        )

        return Response({
            "success": True,
            "message": "Lead Updated Successfully",
            "old_data": old_data,
            "new_data": serializer.data
        })

    except Exception:
        logger.exception("Unexpected error in lead_update_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )


@api_view(['DELETE'])
def lead_delete_api(request, leadid):
    try:
        lead = Lead.objects.filter(pk=leadid).first()

        if not lead:
            return Response(
                {"success": False, "message": "Lead Not Found"},
                status=404
            )

        lead.delete()
        return Response(
            {"success": True, "message": "Lead Deleted Successfully"}
        )

    except Exception:
        logger.exception("Unexpected error in lead_delete_api")
        return Response(
            {"success": False, "message": "Internal server error"},
            status=500
        )